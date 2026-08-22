#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""将华原滤芯滤材规格手写表转为 Excel。"""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins

OUT = "/workspace/华原滤芯滤材规格.xlsx"

thin = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)
header_fill = PatternFill("solid", fgColor="1F4E79")
header_font = Font(name="微软雅黑", bold=True, color="FFFFFF", size=11)
title_font = Font(name="微软雅黑", bold=True, color="1F4E79", size=18)
subtitle_font = Font(name="微软雅黑", size=10, color="5B7A99")
normal_font = Font(name="微软雅黑", size=11)
spec_font = Font(name="微软雅黑", bold=True, size=11, color="1F4E79")
label_font = Font(name="微软雅黑", bold=True, size=10, color="1F4E79")
note_font = Font(name="微软雅黑", size=9, italic=True, color="666666")
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)
alt_fill = PatternFill("solid", fgColor="D6E3F0")
info_fill = PatternFill("solid", fgColor="D6EAF8")
warn_fill = PatternFill("solid", fgColor="FFF3CD")
title_fill = PatternFill("solid", fgColor="E8F0F7")

# 按手写原稿录入。折高取涂改后的最终值。
ROWS = [
    {
        "序号": 1,
        "规格": "YH-1221X-9000",
        "折数": 67,
        "外径": 117.5,
        "折高": 210,
        "折宽": 11,
        "备注": "折高原稿涂改为 210（原写 205.5）",
    },
    {
        "序号": 2,
        "规格": "YH-0932-9000",
        "折数": 64,
        "外径": 84.5,
        "折高": 32,
        "折宽": 14,
        "备注": "折高原稿有涂改，按最终手写值 32 录入",
    },
    {
        "序号": 3,
        "规格": "YH-1345AIH-9000",
        "折数": 108,
        "外径": 174,
        "折高": 448,
        "折宽": 28.5,
        "备注": "",
    },
    {
        "序号": 4,
        "规格": "YH-1534H-9000",
        "折数": 102,
        "外径": 146,
        "折高": 350,
        "折宽": 21.5,
        "备注": "",
    },
    {
        "序号": 5,
        "规格": "YH-1345H-9000",
        "折数": 105,
        "外径": 174,
        "折高": 448,
        "折宽": 28.5,
        "备注": "",
    },
]

HEADERS = ["序号", "规格", "折数", "外径 Φ", "折高", "折宽", "备注"]
COL_KEYS = ["序号", "规格", "折数", "外径", "折高", "折宽", "备注"]
COL_WIDTHS = [8, 22, 10, 12, 10, 10, 42]


def apply_cell(cell, font, alignment, fill=None, number_format=None):
    cell.font = font
    cell.alignment = alignment
    cell.border = thin
    if fill is not None:
        cell.fill = fill
    if number_format is not None:
        cell.number_format = number_format


def write_main_sheet(wb):
    ws = wb.active
    ws.title = "华原滤芯滤材规格"
    last_col = get_column_letter(len(HEADERS))

    ws.merge_cells(f"A1:{last_col}1")
    ws["A1"] = "华原滤芯滤材规格"
    ws["A1"].font = title_font
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws["A1"].fill = title_fill
    ws.row_dimensions[1].height = 32

    ws.merge_cells(f"A2:{last_col}2")
    ws["A2"] = "滤材折叠规格一览（按手写原稿录入，外径/折高/折宽单位：mm）"
    ws["A2"].font = subtitle_font
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws["A2"].fill = info_fill
    ws.row_dimensions[2].height = 20

    header_row = 4
    for c, header in enumerate(HEADERS, 1):
        apply_cell(ws.cell(header_row, c, header), header_font, center, header_fill)
    ws.row_dimensions[header_row].height = 24

    for i, rec in enumerate(ROWS):
        r = header_row + 1 + i
        fill = alt_fill if i % 2 == 1 else None
        ws.row_dimensions[r].height = 26
        for c, key in enumerate(COL_KEYS, 1):
            value = rec[key]
            cell = ws.cell(r, c, value)
            font = spec_font if key == "规格" else normal_font
            align = left_align if key == "备注" else center
            fmt = None
            if key == "规格":
                fmt = "@"
            elif key == "外径":
                fmt = '"Φ"0.0'
            elif key in ("折高", "折数"):
                fmt = "0"
            elif key == "折宽":
                fmt = "0.0"
            apply_cell(cell, font, align, warn_fill if rec["备注"] and key == "备注" else fill, fmt)

    last_data = header_row + len(ROWS)
    note_row = last_data + 2
    ws.merge_cells(f"A{note_row}:{last_col}{note_row}")
    ws[f"A{note_row}"] = (
        "说明：外径栏数字前的 Φ 表示直径；第 1、2 行折高以原稿涂改后的最终值为准。"
        "纸面右上角为跨越速运运单标识，不属于规格内容。"
    )
    ws[f"A{note_row}"].font = note_font
    ws[f"A{note_row}"].alignment = left_align
    ws.row_dimensions[note_row].height = 28

    ws.auto_filter.ref = f"A{header_row}:{last_col}{last_data}"
    ws.freeze_panes = f"A{header_row + 1}"
    for i, width in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(i)].width = width

    ws.print_title_rows = "1:4"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.5, right=0.5, top=0.6, bottom=0.6)
    ws.oddHeader.left.text = "华原滤芯滤材规格"
    ws.oddFooter.right.text = "第 &P 页 / 共 &N 页"
    ws.oddFooter.left.text = "按手写原稿录入"
    return ws


def write_readme_sheet(wb):
    ws = wb.create_sheet("录入说明")
    ws.merge_cells("A1:B1")
    ws["A1"] = "录入说明"
    ws["A1"].font = title_font
    ws["A1"].fill = title_fill
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 28

    ws["A3"] = "项目"
    ws["B3"] = "内容"
    apply_cell(ws["A3"], header_font, center, header_fill)
    apply_cell(ws["B3"], header_font, center, header_fill)
    ws.row_dimensions[3].height = 22

    rows = [
        ("表格名称", "华原滤芯滤材规格"),
        ("来源", "手写规格表照片（表格竖排；纸面右上角有跨越速运标识）"),
        ("列对应", "规格、折数、外径、折高、折宽（与原稿列名一致）；本表另加序号、备注便于查阅"),
        ("单位", "折数：个；外径 / 折高 / 折宽：mm"),
        ("外径符号", "原稿数字前写 Φ，表示直径；Excel 中用数字格式显示为 Φxxx.x"),
        ("第1行折高", "YH-1221X-9000 原稿先写 205.5，后涂改为 210，按 210 录入"),
        ("第2行折高", "YH-0932-9000 原稿折高有涂改，按最终手写值 32 录入"),
        ("其余行", "YH-1345AIH-9000、YH-1534H-9000、YH-1345H-9000 按原稿清晰数字录入"),
        ("筛选", "主表表头已开启筛选，可按规格或尺寸过滤"),
    ]
    for i, (key, value) in enumerate(rows):
        r = 4 + i
        apply_cell(ws.cell(r, 1, key), label_font, center, info_fill)
        apply_cell(ws.cell(r, 2, value), normal_font, left_align)
        ws.row_dimensions[r].height = 28

    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 88
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    return ws


def main():
    wb = Workbook()
    write_main_sheet(wb)
    write_readme_sheet(wb)
    wb.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
