#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""将样品核对纸表转为 Excel：印刷栏原样录入，手写内容写入「问题描述」。"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import FormulaRule
from openpyxl.chart import PieChart, Reference
from openpyxl.chart.label import DataLabelList
from datetime import date

OUT = "/workspace/样品订单核对表_HD26YDKCYY001.xlsx"

# ---------- styles ----------
thin = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)
header_fill = PatternFill("solid", fgColor="1F4E79")
header_font = Font(name="微软雅黑", bold=True, color="FFFFFF", size=11)
title_font = Font(name="微软雅黑", bold=True, color="1F4E79", size=16)
section_font = Font(name="微软雅黑", bold=True, color="1F4E79", size=12)
normal_font = Font(name="微软雅黑", size=10)
note_font = Font(name="微软雅黑", size=9, italic=True, color="666666")
label_font = Font(name="微软雅黑", bold=True, size=10, color="1F4E79")
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)
alt_fill = PatternFill("solid", fgColor="D6E3F0")
info_fill = PatternFill("solid", fgColor="D6EAF8")
warn_fill = PatternFill("solid", fgColor="FFF3CD")
issue_fill = PatternFill("solid", fgColor="FCE4D6")
fail_fill = PatternFill("solid", fgColor="F8D7DA")
ok_fill = PatternFill("solid", fgColor="D4EDDA")
unclear_fill = PatternFill("solid", fgColor="E2E3E5")


# 印刷表字段 + 手写「问题描述」
# 核对标记：纸表手写勾/叉；问题描述：手写文字（不含勾叉）
ROWS = [
    {
        "序号": 1,
        "样品订单号": "HD26YDKCYY001",
        "样品日期": date(2026, 3, 5),
        "OEM号/客户号": "186-1012240",
        "HD号": "DJ8890-DKCYY",
        "客户签收样品数": 11,
        "回取样品数量": 7,
        "核对标记": "✓",
        "问题描述": "盖板不同",
        "墨迹": "黑",
        "状态": "有问题",
    },
    {
        "序号": 2,
        "样品订单号": "HD26YDKCYY001",
        "样品日期": date(2026, 3, 5),
        "OEM号/客户号": "A50000-1105350",
        "HD号": "DC4173-DKCYY",
        "客户签收样品数": 11,
        "回取样品数量": 7,
        "核对标记": "✓",
        "问题描述": "少带一个紧固卡，外面密封圈不同",
        "墨迹": "黑",
        "状态": "有问题",
    },
    {
        "序号": 3,
        "样品订单号": "HD26YDKCYY001",
        "样品日期": date(2026, 3, 5),
        "OEM号/客户号": "129907-55800",
        "HD号": "DC4174-DKCYY",
        "客户签收样品数": 11,
        "回取样品数量": 7,
        "核对标记": "✓",
        "问题描述": "",
        "墨迹": "黑",
        "状态": "已核对",
    },
    {
        "序号": 4,
        "样品订单号": "HD26YDKCYY001",
        "样品日期": date(2026, 3, 5),
        "OEM号/客户号": "A50900-1105140",
        "HD号": "DC4176-DKCYY",
        "客户签收样品数": 11,
        "回取样品数量": 7,
        "核对标记": "✓",
        "问题描述": "",
        "墨迹": "黑（加粗）",
        "状态": "已核对",
    },
    {
        "序号": 5,
        "样品订单号": "HD26YDKCYY001",
        "样品日期": date(2026, 3, 5),
        "OEM号/客户号": "26561117",
        "HD号": "DC6293I-DKCYY",
        "客户签收样品数": 11,
        "回取样品数量": 7,
        "核对标记": "✓",
        "问题描述": "",
        "墨迹": "黑（加粗）",
        "状态": "已核对",
    },
    {
        "序号": 6,
        "样品订单号": "HD26YDKCYY001",
        "样品日期": date(2026, 3, 5),
        "OEM号/客户号": "HH16432430",
        "HD号": "DJ8784-DKCYY",
        "客户签收样品数": 11,
        "回取样品数量": 7,
        "核对标记": "✗",
        "问题描述": "盖板不同",
        "墨迹": "黑",
        "状态": "不合格",
    },
    {
        "序号": 7,
        "样品订单号": "HD26YDKCYY001",
        "样品日期": date(2026, 3, 5),
        "OEM号/客户号": "HH1C0-32430",
        "HD号": "DJ8892-DKCYY",
        "客户签收样品数": 11,
        "回取样品数量": 7,
        "核对标记": "✓",
        "问题描述": "",
        "墨迹": "黑（加粗）",
        "状态": "已核对",
    },
    {
        "序号": 8,
        "样品订单号": "HD26YDKCYY001",
        "样品日期": date(2026, 3, 5),
        "OEM号/客户号": "119305-35160",
        "HD号": "DJ8523I-DKCYY",
        "客户签收样品数": 11,
        "回取样品数量": 7,
        "核对标记": "✓",
        "问题描述": "紧固件卡扣不同",
        "墨迹": "红",
        "状态": "有问题",
    },
    {
        "序号": 9,
        "样品订单号": "HD26YDKCYY001",
        "样品日期": date(2026, 3, 5),
        "OEM号/客户号": "HH16432430",
        "HD号": "DJ8795-DKCYY",
        "客户签收样品数": 11,
        "回取样品数量": 7,
        "核对标记": "?",
        "问题描述": "（手写潦草，无法准确辨认）",
        "墨迹": "黑",
        "状态": "待确认",
    },
    {
        "序号": 10,
        "样品订单号": "HD26YDKCYY001",
        "样品日期": date(2026, 3, 5),
        "OEM号/客户号": "FS20019",
        "HD号": "DC4048I-DKCYY",
        "客户签收样品数": 11,
        "回取样品数量": 7,
        "核对标记": "✓",
        "问题描述": "边框底端号",
        "墨迹": "红",
        "状态": "有问题",
    },
    {
        "序号": 11,
        "样品订单号": "HD26YDKCYY001",
        "样品日期": date(2026, 3, 5),
        "OEM号/客户号": "5266016",
        "HD号": "DJ7267-DKCYY",
        "客户签收样品数": 11,
        "回取样品数量": 7,
        "核对标记": "✓",
        "问题描述": "",
        "墨迹": "黑（加粗）",
        "状态": "已核对",
    },
    {
        "序号": 12,
        "样品订单号": "HD26YDKCYY001",
        "样品日期": date(2026, 3, 5),
        "OEM号/客户号": "HH166-43560",
        "HD号": "DC4175-DKCYY",
        "客户签收样品数": 11,
        "回取样品数量": 7,
        "核对标记": "✓",
        "问题描述": "",
        "墨迹": "黑（加粗）",
        "状态": "已核对",
    },
    {
        "序号": 13,
        "样品订单号": "HD26YDKCYY001",
        "样品日期": date(2026, 3, 5),
        "OEM号/客户号": "FS36209",
        "HD号": "DC6286III-DKCYY",
        "客户签收样品数": 11,
        "回取样品数量": 7,
        "核对标记": "✓",
        "问题描述": "边孔尺寸最终确认",
        "墨迹": "红",
        "状态": "有问题",
    },
]

HEADERS = [
    "序号",
    "样品订单号",
    "样品日期",
    "OEM号/客户号",
    "HD号",
    "客户签收样品数",
    "回取样品数量",
    "核对标记",
    "问题描述",
    "墨迹",
    "状态",
]


def apply_border_font(cell, font=None, align=None, fill=None):
    cell.border = thin
    cell.font = font or normal_font
    cell.alignment = align or center
    if fill is not None:
        cell.fill = fill


def status_fill(status):
    return {
        "有问题": issue_fill,
        "不合格": fail_fill,
        "已核对": ok_fill,
        "待确认": unclear_fill,
    }.get(status)


def write_main_sheet(wb):
    ws = wb.active
    ws.title = "样品核对表"

    ws.merge_cells("A1:K1")
    ws["A1"] = "样品订单核对表（纸表录入）"
    ws["A1"].font = title_font
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 28

    meta = [
        ("A2", "样品订单号", "B2", "HD26YDKCYY001"),
        ("C2", "样品日期", "D2", date(2026, 3, 5)),
        ("E2", "客户签收样品数", "F2", 11),
        ("G2", "回取样品数量", "H2", 7),
        ("I2", "行数", "J2", 13),
    ]
    for lab, lab_val, val, val_val in meta:
        ws[lab] = lab_val
        ws[lab].font = label_font
        ws[lab].fill = info_fill
        ws[lab].alignment = center
        ws[lab].border = thin
        ws[val] = val_val
        ws[val].font = normal_font
        ws[val].alignment = center
        ws[val].border = thin
        if isinstance(val_val, date):
            ws[val].number_format = "YYYY/M/D"
    ws.merge_cells("J2:K2")
    ws["K2"].border = thin
    ws.row_dimensions[2].height = 22

    ws.merge_cells("A3:K3")
    ws["A3"] = (
        "说明：印刷栏按原表录入；纸表手写勾/叉记入「核对标记」，手写文字记入「问题描述」。"
        "第9行手写潦草无法准确辨认；表左侧有「鲁」字圆形印章。红笔备注见「墨迹」列。"
    )
    ws["A3"].font = note_font
    ws["A3"].alignment = left_align
    ws["A3"].fill = warn_fill
    ws.row_dimensions[3].height = 32

    header_row = 4
    for col, h in enumerate(HEADERS, 1):
        cell = ws.cell(header_row, col, h)
        apply_border_font(cell, header_font, center, header_fill)
    ws.row_dimensions[header_row].height = 24
    ws.auto_filter.ref = f"A{header_row}:K{header_row + len(ROWS)}"
    ws.freeze_panes = "A5"

    for i, rec in enumerate(ROWS):
        r = header_row + 1 + i
        fill = status_fill(rec["状态"])
        if fill is None and i % 2 == 1:
            fill = alt_fill
        values = [rec[h] for h in HEADERS]
        for c, val in enumerate(values, 1):
            cell = ws.cell(r, c, val)
            font = normal_font
            align = center
            if HEADERS[c - 1] == "问题描述":
                align = left_align
                if rec["状态"] in ("有问题", "不合格"):
                    font = Font(name="微软雅黑", size=10, color="C0392B", bold=True)
                elif rec["状态"] == "待确认":
                    font = Font(name="微软雅黑", size=10, italic=True, color="666666")
            if HEADERS[c - 1] == "核对标记":
                if val == "✗":
                    font = Font(name="微软雅黑", size=12, color="C0392B", bold=True)
                elif val == "✓":
                    font = Font(name="微软雅黑", size=12, color="196F3D", bold=True)
                elif val == "?":
                    font = Font(name="微软雅黑", size=12, color="6C757D", bold=True)
            if HEADERS[c - 1] == "OEM号/客户号" or HEADERS[c - 1] == "HD号":
                cell.number_format = "@"
            if HEADERS[c - 1] == "样品日期":
                cell.number_format = "YYYY/M/D"
            apply_border_font(cell, font, align, fill)
        ws.row_dimensions[r].height = 22 if not rec["问题描述"] or len(rec["问题描述"]) < 18 else 32

    widths = [8, 18, 12, 20, 18, 16, 14, 12, 36, 12, 12]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.print_title_rows = "1:4"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.oddHeader.left.text = "样品订单 HD26YDKCYY001"
    ws.oddFooter.right.text = "第 &P 页 / 共 &N 页"

    last = header_row + len(ROWS)
    ws.conditional_formatting.add(
        f"K5:K{last}",
        FormulaRule(formula=['K5="不合格"'], fill=fail_fill),
    )
    return ws


def write_issue_sheet(wb):
    ws = wb.create_sheet("问题汇总")
    ws.merge_cells("A1:G1")
    ws["A1"] = "问题汇总（仅含手写问题描述的行）"
    ws["A1"].font = title_font
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 28

    ws.merge_cells("A2:G2")
    ws["A2"] = "从「样品核对表」筛选：问题描述非空的记录，便于跟进盖板、密封圈、卡扣、边框、边孔等差异。"
    ws["A2"].font = note_font
    ws["A2"].fill = warn_fill
    ws["A2"].alignment = left_align

    headers = ["序号", "OEM号/客户号", "HD号", "核对标记", "问题描述", "墨迹", "状态"]
    for c, h in enumerate(headers, 1):
        cell = ws.cell(4, c, h)
        apply_border_font(cell, header_font, center, header_fill)
    ws.row_dimensions[4].height = 24

    issues = [r for r in ROWS if r["问题描述"]]
    for i, rec in enumerate(issues):
        r = 5 + i
        fill = status_fill(rec["状态"])
        vals = [rec[h] for h in headers]
        for c, val in enumerate(vals, 1):
            cell = ws.cell(r, c, val)
            font = normal_font
            align = left_align if headers[c - 1] == "问题描述" else center
            if headers[c - 1] == "问题描述" and rec["状态"] != "待确认":
                font = Font(name="微软雅黑", size=10, color="C0392B", bold=True)
            if headers[c - 1] == "OEM号/客户号" or headers[c - 1] == "HD号":
                cell.number_format = "@"
            apply_border_font(cell, font, align, fill)
        ws.row_dimensions[r].height = 28

    last = 4 + len(issues)
    ws.auto_filter.ref = f"A4:G{last}"
    ws.freeze_panes = "A5"
    for i, w in enumerate([8, 20, 18, 12, 42, 12, 12], 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # 统计
    ws["A13"] = "状态统计"
    ws["A13"].font = section_font
    ws.merge_cells("A13:C13")

    stats_headers = ["状态", "数量", "说明"]
    for c, h in enumerate(stats_headers, 1):
        cell = ws.cell(14, c, h)
        apply_border_font(cell, header_font, center, header_fill)

    stats = [
        ("已核对", sum(1 for r in ROWS if r["状态"] == "已核对"), "仅有勾选，无文字问题"),
        ("有问题", sum(1 for r in ROWS if r["状态"] == "有问题"), "已勾选并写明差异/待确认事项"),
        ("不合格", sum(1 for r in ROWS if r["状态"] == "不合格"), "手写叉号，盖板不同"),
        ("待确认", sum(1 for r in ROWS if r["状态"] == "待确认"), "手写潦草无法辨认"),
        ("合计", len(ROWS), "本表共 13 条样品"),
    ]
    for i, (st, n, desc) in enumerate(stats):
        r = 15 + i
        fill = status_fill(st) if st != "合计" else info_fill
        font = Font(name="微软雅黑", bold=True, size=10) if st == "合计" else normal_font
        for c, val in enumerate((st, n, desc), 1):
            cell = ws.cell(r, c, val)
            apply_border_font(cell, font, center if c < 3 else left_align, fill)

    pie = PieChart()
    pie.title = "核对状态分布"
    labels = Reference(ws, min_col=1, min_row=15, max_row=18)
    data = Reference(ws, min_col=2, min_row=14, max_row=18)
    pie.add_data(data, titles_from_data=True)
    pie.set_categories(labels)
    pie.dataLabels = DataLabelList()
    pie.dataLabels.showPercent = True
    pie.dataLabels.showVal = True
    pie.dataLabels.showCatName = False
    pie.width = 12
    pie.height = 8
    ws.add_chart(pie, "E13")

    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    return ws


def write_readme_sheet(wb):
    ws = wb.create_sheet("录入说明")
    ws.merge_cells("A1:B1")
    ws["A1"] = "录入说明"
    ws["A1"].font = title_font
    ws.row_dimensions[1].height = 28

    rows = [
        ("来源", "纸质样品核对表照片（表格竖排，照片横拍）"),
        ("样品订单号", "HD26YDKCYY001（13 行相同）"),
        ("样品日期", "2026/3/5（13 行相同）"),
        ("客户签收样品数", "11（表内各行填写）"),
        ("回取样品数量", "7（表内各行填写）"),
        ("印刷列", "序号、样品订单号、样品日期、OEM号/客户号、HD号、客户签收样品数、回取样品数量"),
        ("手写列处理", "原表该列印刷表头亦为「OEM号/客户号」，但单元格为手写勾/叉与备注；按要求将该列命名为「问题描述」，并拆出「核对标记」"),
        ("问题描述", "仅录入手写汉字内容，不含勾叉"),
        ("核对标记", "✓ 已勾选；✗ 叉号（第6行）；? 无法辨认（第9行）"),
        ("墨迹", "黑 / 黑（加粗）/ 红，对应纸表笔迹颜色"),
        ("印章", "表左侧近第13行有圆形「鲁」字印章"),
        ("辨认说明", "第9行手写符号潦草，问题描述标注为无法准确辨认，请以原件复核"),
        ("HD号后缀", "原表中 I / III 按罗马数字型号后缀录入，如 DC6293I-DKCYY、DC6286III-DKCYY"),
        ("筛选", "表头已开启筛选；「问题汇总」仅列出有手写文字的行"),
    ]
    ws["A3"] = "项目"
    ws["B3"] = "内容"
    apply_border_font(ws["A3"], header_font, center, header_fill)
    apply_border_font(ws["B3"], header_font, center, header_fill)
    for i, (k, v) in enumerate(rows):
        r = 4 + i
        apply_border_font(ws.cell(r, 1, k), label_font, center, info_fill)
        apply_border_font(ws.cell(r, 2, v), normal_font, left_align)
        ws.row_dimensions[r].height = 28
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 88
    ws.row_dimensions[3].height = 22
    return ws


def main():
    wb = Workbook()
    write_main_sheet(wb)
    write_issue_sheet(wb)
    write_readme_sheet(wb)
    wb.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
